#!/usr/bin/env python3
"""Descarga y unifica localmente los archivos mensuales de emergencias del ECU 911.

El programa trabaja por streaming: descarga y procesa un solo mes a la vez. Conserva
los siete campos originales y agrega tres columnas de trazabilidad. El resultado usa
el mismo separador ``;`` de los archivos fuente.
"""

from __future__ import annotations

import argparse
import codecs
import csv
import os
import re
import shutil
import sys
import tempfile
import time
import unicodedata
import urllib.error
import urllib.request
from datetime import date, datetime, time as datetime_time
from pathlib import Path
from urllib.parse import unquote, urlparse


CATALOG_URL = (
    "https://www.datosabiertos.gob.ec/api/3/action/"
    "package_show?id=base-de-emergencias"
)
PERIOD_PATTERN = re.compile(r"(?<!\d)(20\d{4})(?!\d)")
SOURCE_COLUMNS = (
    "Fecha",
    "provincia",
    "Canton",
    "Cod_Parroquia",
    "Parroquia",
    "Servicio",
    "Subtipo",
)
OUTPUT_COLUMNS = SOURCE_COLUMNS + (
    "_periodo",
    "_archivo_origen",
    "_fila_archivo",
)
USER_AGENT = "unificador-local-ecu911/1.0"


def parse_args() -> argparse.Namespace:
    laboratorios = Path(__file__).resolve().parent
    parser = argparse.ArgumentParser(
        description="Crea un CSV crudo con todos los meses publicados por el ECU 911."
    )
    parser.add_argument(
        "--desde",
        help="Primer periodo, con formato AAAAMM. Por defecto usa el primero disponible.",
    )
    parser.add_argument(
        "--hasta",
        help="Ultimo periodo, con formato AAAAMM. Por defecto usa el ultimo disponible.",
    )
    parser.add_argument(
        "--salida",
        type=Path,
        default=laboratorios
        / "semana-08-sesion-02"
        / "data"
        / "ecu911_emergencias_raw.csv",
        help="Ruta del CSV principal.",
    )
    parser.add_argument(
        "--enlace",
        type=Path,
        default=laboratorios
        / "semana-09-sesion-01"
        / "data"
        / "ecu911_emergencias_raw.csv",
        help="Segunda ruta que apuntara al mismo archivo fisico.",
    )
    parser.add_argument(
        "--forzar",
        action="store_true",
        help="Reemplaza una salida existente solamente despues de terminar la nueva.",
    )
    return parser.parse_args()


def validate_period(value: str | None, option: str) -> None:
    if value is None:
        return
    if not re.fullmatch(r"20\d{4}", value):
        raise ValueError(f"{option} debe tener el formato AAAAMM: {value!r}")
    month = int(value[4:])
    if month < 1 or month > 12:
        raise ValueError(f"{option} contiene un mes invalido: {value!r}")


def request_json(url: str, retries: int = 4) -> dict:
    last_error: Exception | None = None
    for attempt in range(1, retries + 1):
        try:
            request = urllib.request.Request(url, headers={"User-Agent": USER_AGENT})
            with urllib.request.urlopen(request, timeout=120) as response:
                import json

                return json.load(response)
        except (OSError, urllib.error.URLError) as error:
            last_error = error
            if attempt == retries:
                break
            wait_seconds = 2**attempt
            print(
                f"  Catalogo no disponible (intento {attempt}/{retries}); "
                f"reintentando en {wait_seconds}s...",
                flush=True,
            )
            time.sleep(wait_seconds)
    raise RuntimeError("No se pudo consultar el catalogo del ECU 911") from last_error


def catalog_resources(start: str | None, end: str | None) -> list[dict]:
    payload = request_json(CATALOG_URL)
    if not payload.get("success"):
        raise RuntimeError("La API CKAN respondio sin exito")

    by_period: dict[str, dict] = {}
    for resource in payload["result"]["resources"]:
        searchable = f"{resource.get('name', '')} {resource.get('url', '')}"
        match = PERIOD_PATTERN.search(searchable)
        if not match:
            continue
        period = match.group(1)
        if start and period < start:
            continue
        if end and period > end:
            continue
        if period in by_period:
            raise RuntimeError(f"El catalogo contiene dos recursos para {period}")
        by_period[period] = {**resource, "period": period}

    resources = [by_period[key] for key in sorted(by_period)]
    if not resources:
        raise RuntimeError("No se encontraron recursos para el rango solicitado")
    validate_contiguous_periods([resource["period"] for resource in resources])
    return resources


def next_period(period: str) -> str:
    year = int(period[:4])
    month = int(period[4:])
    if month == 12:
        return f"{year + 1}01"
    return f"{year}{month + 1:02d}"


def validate_contiguous_periods(periods: list[str]) -> None:
    for previous, current in zip(periods, periods[1:]):
        expected = next_period(previous)
        if current != expected:
            raise RuntimeError(
                f"Falta al menos un mes en el catalogo: despues de {previous} "
                f"se esperaba {expected}, pero aparece {current}"
            )


def source_filename(resource: dict) -> str:
    filename = Path(unquote(urlparse(resource["url"]).path)).name
    return filename or f"ecu911_{resource['period']}"


def download(resource: dict, destination: Path, retries: int = 4) -> None:
    expected_size = resource.get("size")
    try:
        expected_size = int(expected_size) if expected_size is not None else None
    except (TypeError, ValueError):
        expected_size = None

    last_error: Exception | None = None
    for attempt in range(1, retries + 1):
        try:
            request = urllib.request.Request(
                resource["url"], headers={"User-Agent": USER_AGENT}
            )
            with urllib.request.urlopen(request, timeout=300) as response:
                with destination.open("wb") as output:
                    shutil.copyfileobj(response, output, length=1024 * 1024)
            actual_size = destination.stat().st_size
            if expected_size and actual_size != expected_size:
                raise IOError(
                    f"tamano incompleto: {actual_size:,} de {expected_size:,} bytes"
                )
            return
        except (OSError, urllib.error.URLError) as error:
            last_error = error
            destination.unlink(missing_ok=True)
            if attempt == retries:
                break
            wait_seconds = 2**attempt
            print(
                f"  Descarga fallida (intento {attempt}/{retries}: {error}); "
                f"reintentando en {wait_seconds}s...",
                flush=True,
            )
            time.sleep(wait_seconds)
    raise RuntimeError(f"No se pudo descargar {resource['url']}") from last_error


def normalize_header(value: object) -> str:
    text = "" if value is None else str(value)
    text = unicodedata.normalize("NFKD", text.strip().lstrip("\ufeff"))
    text = "".join(character for character in text if not unicodedata.combining(character))
    return re.sub(r"[^a-z0-9]", "", text.lower())


def column_indices(header: list[object]) -> list[int]:
    expected = [normalize_header(column) for column in SOURCE_COLUMNS]
    normalized = [normalize_header(column) for column in header]
    unexpected = [
        str(value)
        for value, key in zip(header, normalized)
        if key and key not in expected
    ]
    if unexpected:
        raise ValueError(f"Columnas no esperadas en la fuente: {unexpected}")

    indices: list[int] = []
    for column, key in zip(SOURCE_COLUMNS, expected):
        matches = [index for index, candidate in enumerate(normalized) if candidate == key]
        if len(matches) != 1:
            raise ValueError(
                f"La columna {column!r} aparece {len(matches)} veces en {header!r}"
            )
        indices.append(matches[0])
    return indices


def detect_csv_encoding(path: Path) -> str:
    decoder = codecs.getincrementaldecoder("utf-8-sig")("strict")
    try:
        with path.open("rb") as source:
            while chunk := source.read(1024 * 1024):
                decoder.decode(chunk)
            decoder.decode(b"", final=True)
        return "utf-8-sig"
    except UnicodeDecodeError:
        return "cp1252"


def write_csv_resource(
    path: Path,
    resource: dict,
    writer: csv.writer,
) -> int:
    encoding = detect_csv_encoding(path)
    period = resource["period"]
    filename = source_filename(resource)
    rows_written = 0

    with path.open("r", encoding=encoding, newline="") as source:
        reader = csv.reader(source, delimiter=";")
        try:
            header = next(reader)
        except StopIteration as error:
            raise ValueError(f"El archivo {filename} esta vacio") from error
        indices = column_indices(header)

        for source_row_number, row in enumerate(reader, start=2):
            if not row or not any(value.strip() for value in row):
                continue
            if max(indices) >= len(row):
                raise ValueError(
                    f"{filename}, fila {source_row_number}: faltan columnas ({row!r})"
                )
            extra_values = row[len(header) :]
            if any(value.strip() for value in extra_values):
                raise ValueError(
                    f"{filename}, fila {source_row_number}: sobran valores ({row!r})"
                )
            writer.writerow(
                [row[index] for index in indices]
                + [period, filename, source_row_number]
            )
            rows_written += 1
    return rows_written


def excel_value(value: object, epoch: datetime) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date, datetime_time)):
        from openpyxl.utils.datetime import to_excel

        serial = to_excel(value, epoch=epoch)
        return format(serial, ".15g")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def write_xlsx_resource(
    path: Path,
    resource: dict,
    writer: csv.writer,
) -> int:
    try:
        from openpyxl import load_workbook
    except ModuleNotFoundError as error:
        raise RuntimeError(
            "Se necesita openpyxl para leer los meses publicados como Excel"
        ) from error

    period = resource["period"]
    filename = source_filename(resource)
    # Se pasa un flujo y no la ruta porque algunos recursos son XLSX aunque su URL
    # termine en .csv; openpyxl valida la extension cuando recibe un nombre.
    with path.open("rb") as binary_source:
        workbook = load_workbook(binary_source, read_only=True, data_only=True)
        try:
            worksheet = workbook.worksheets[0]
            iterator = worksheet.iter_rows(values_only=True)

            header: list[object] | None = None
            header_row_number = 0
            for header_row_number, candidate in enumerate(iterator, start=1):
                if any(value not in (None, "") for value in candidate):
                    header = list(candidate)
                    break
            if header is None:
                raise ValueError(f"El archivo {filename} no contiene filas")
            indices = column_indices(header)

            rows_written = 0
            for source_row_number, row in enumerate(
                iterator, start=header_row_number + 1
            ):
                if not row or not any(value not in (None, "") for value in row):
                    continue
                if max(indices) >= len(row):
                    raise ValueError(
                        f"{filename}, fila {source_row_number}: faltan columnas"
                    )
                extra_values = row[len(header) :]
                if any(value not in (None, "") for value in extra_values):
                    raise ValueError(
                        f"{filename}, fila {source_row_number}: sobran valores"
                    )
                writer.writerow(
                    [excel_value(row[index], workbook.epoch) for index in indices]
                    + [period, filename, source_row_number]
                )
                rows_written += 1
            return rows_written
        finally:
            workbook.close()


def is_xlsx(path: Path) -> bool:
    with path.open("rb") as source:
        return source.read(4) == b"PK\x03\x04"


def create_hard_link(source: Path, link: Path) -> None:
    link.parent.mkdir(parents=True, exist_ok=True)
    if link.exists():
        link.unlink()
    try:
        os.link(source, link)
    except OSError as error:
        raise RuntimeError(
            f"El CSV principal se creo, pero Windows no permitio enlazarlo en {link}. "
            "No se hizo una copia para evitar duplicar varios GB."
        ) from error


def main() -> int:
    args = parse_args()
    validate_period(args.desde, "--desde")
    validate_period(args.hasta, "--hasta")
    if args.desde and args.hasta and args.desde > args.hasta:
        raise ValueError("--desde no puede ser posterior a --hasta")

    output = args.salida.resolve()
    link = args.enlace.resolve()
    if output == link:
        raise ValueError("--salida y --enlace deben ser rutas diferentes")
    if output.exists() and not args.forzar:
        raise FileExistsError(
            f"La salida ya existe: {output}. Usa --forzar para regenerarla."
        )

    resources = catalog_resources(args.desde, args.hasta)
    print(
        f"Se unificaran {len(resources)} meses: "
        f"{resources[0]['period']} a {resources[-1]['period']}",
        flush=True,
    )

    output.parent.mkdir(parents=True, exist_ok=True)
    partial = output.with_name(output.name + ".partial")
    partial.unlink(missing_ok=True)
    total_rows = 0

    try:
        with tempfile.TemporaryDirectory(prefix="ecu911_") as temporary_directory:
            downloaded = Path(temporary_directory) / "source"
            with partial.open("w", encoding="utf-8", newline="") as destination:
                writer = csv.writer(
                    destination,
                    delimiter=";",
                    quotechar='"',
                    quoting=csv.QUOTE_MINIMAL,
                    lineterminator="\n",
                )
                writer.writerow(OUTPUT_COLUMNS)

                for position, resource in enumerate(resources, start=1):
                    period = resource["period"]
                    filename = source_filename(resource)
                    print(
                        f"[{position:02d}/{len(resources):02d}] {period}: "
                        f"descargando {filename}...",
                        flush=True,
                    )
                    download(resource, downloaded)
                    if is_xlsx(downloaded):
                        count = write_xlsx_resource(downloaded, resource, writer)
                        source_type = "XLSX"
                    else:
                        count = write_csv_resource(downloaded, resource, writer)
                        source_type = "CSV"
                    if count == 0:
                        raise ValueError(f"El periodo {period} no contiene datos")
                    total_rows += count
                    print(
                        f"  {count:,} filas ({source_type}); "
                        f"acumulado: {total_rows:,}",
                        flush=True,
                    )
                    downloaded.unlink(missing_ok=True)

        os.replace(partial, output)
        create_hard_link(output, link)
    except BaseException:
        partial.unlink(missing_ok=True)
        raise

    if not os.path.samefile(output, link):
        raise RuntimeError("Las dos rutas no apuntan al mismo archivo fisico")

    print("\nUnificacion terminada.", flush=True)
    print(f"Filas de datos: {total_rows:,}", flush=True)
    print(f"Tamano: {output.stat().st_size / (1024**3):.2f} GiB", flush=True)
    print(f"CSV principal: {output}", flush=True)
    print(f"Enlace semana 9: {link}", flush=True)
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except KeyboardInterrupt:
        print("\nProceso cancelado; se elimino la salida parcial.", file=sys.stderr)
        raise SystemExit(130)
    except Exception as error:
        print(f"\nERROR: {error}", file=sys.stderr)
        raise SystemExit(1)
